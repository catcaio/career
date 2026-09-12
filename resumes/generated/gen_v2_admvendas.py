#!/usr/bin/env python3
"""Curriculo 915050 V2 — CAUSA CERTA / Analista Adm Vendas.
Motor: skill resume-document-design (tokens INK/MUTED/ACCENT/RULE, escala V2).
Mesmos fatos do V1 (spec-admvendas-causacerta-v1.json); sem metrica de faturamento.
Gera PDF (reportlab) + DOCX (python-docx) + spec JSON V2.
Uso: python3 gen_v2_admvendas.py
"""
import json, os
from copy import deepcopy

OUT_DIR = os.path.dirname(os.path.abspath(__file__))
PDF_OUT = os.path.join(OUT_DIR, "Rafa_Nascimento_Curriculo_CausaCerta_AnalistaAdmVendas_V2.pdf")
DOCX_OUT = os.path.join(OUT_DIR, "Rafa_Nascimento_Curriculo_CausaCerta_AnalistaAdmVendas_V2.docx")
SPEC_OUT = os.path.join(OUT_DIR, "spec-admvendas-causacerta-v2.json")

INK, MUTED, ACCENT, RULE = "23272F", "5A6270", "1F5C4E", "C9CFD6"

CONTENT = {
    "name": "CAIO RAFAEL BARROS DO NASCIMENTO",
    "role": "Analista Administrativo de Vendas  |  Rotinas Administrativas  |  Logística e Suporte Operacional",
    "contact": "Nova Palhoça — Palhoça/SC  |  (48) 99153-2952  |  cr_bn20@hotmail.com  |  linkedin.com/in/rafael-brasiliae",
    "sections": [
        {"title": "RESUMO PROFISSIONAL", "body": [
            "Profissional administrativo e operacional com mais de 5 anos de atuação em rotinas administrativas de escritório, suporte operacional, logística de compras e suprimentos. Experiência com documentos, planilhas, sistemas, contratos, gestão de fretes, Incoterms CIF e FOB, ERP, CRM e licitações eletrônicas, com acompanhamento ponta a ponta do recebimento na fábrica à entrega final. Perfil organizado, atento a detalhes e orientado à consolidação de processos confiáveis."
        ]},
        {"title": "COMPETÊNCIAS-CHAVE", "bullets": [
            "Rotinas administrativas e suporte operacional: documentos, planilhas, sistemas, registros, bases e acompanhamento de demandas internas",
            "Logística de compras e suprimentos: gestão de fretes, Incoterms CIF e FOB e acompanhamento do recebimento na fábrica à entrega final",
            "Contratos e documentação: conferência, organização, tabelas, condições de transporte, status e prazos",
            "Sistemas: LINVIX (Tecmater Master), Tiny/Olist, RD Station, Loja Integrada, WordPress/WooCommerce e e-mail marketing",
            "Licitações eletrônicas e pregões: participação em nome da empresa, negociação com fornecedores e acompanhamento dos processos"
        ]},
        {"title": "EXPERIÊNCIA PROFISSIONAL", "jobs": [
            {"company": "LOJACOND EQUIPAMENTOS PARA CONDOMÍNIOS — Palhoça/SC  |  Mar/2021 – Abr/2026",
             "roles": [
                {"role": "Gestor de E-commerce e Operações Digitais  |  Out/2024 – Abr/2026", "bullets": [
                    "Evolução interna da função administrativa e operacional para a gestão integrada de e-commerce, loja online e marketplaces",
                    "Gestão de vendas, catálogo, pedidos, estoque, preços, anúncios e indicadores; administração de aproximadamente 2 mil SKUs e base de 7 mil clientes",
                    "Experiência de 1 ano na Loja Integrada, liderando a migração de site catálogo/orçamento para e-commerce transacional, seguida de nova migração técnica para WordPress com WooCommerce",
                    "Integração entre e-commerce, ERP (Olist/Tiny), pagamentos, frete, estoque e canais de atendimento; acompanhamento do fluxo completo do pedido",
                    "Gestão de CRM e e-mail marketing; gestão da loja oficial no Mercado Livre e melhoria contínua da operação"
                ]},
                {"role": "Assistente Administrativo e Comercial  |  Mar/2021 – Out/2024", "bullets": [
                    "Ingresso como estagiário durante a graduação em Administração, efetivado em CLT após 3 meses, com crescimento interno e ampliação contínua de responsabilidades",
                    "Execução e organização de rotinas administrativas de escritório e suporte operacional: documentos, planilhas, sistemas, registros, bases e acompanhamento de demandas internas",
                    "Atuação técnica em logística de compras e suprimentos, gestão de fretes e Incoterms CIF e FOB, com acompanhamento ponta a ponta do recebimento na fábrica à entrega final",
                    "Uso do ERP LINVIX, da Tecmater Master, durante 3 anos, com posterior migração de processos para o Tiny, da Olist; uso do CRM da RD Station",
                    "No último ano da atuação na empresa, participação em licitações em nome da empresa e domínio dos processos de licitações eletrônicas, incluindo pregões vencidos por meio da negociação com a lista de fornecedores",
                    "Negociação e manutenção de contratos, tabelas e condições com transportadoras e parceiros; apoio pontual a atendimento, orçamento e demandas comerciais"
                ]},
             ]},
            {"company": "FLEX CONTACT CENTER — Florianópolis/SC  |  Fev/2018 – Dez/2020",
             "roles": [
                {"role": None, "bullets": [
                    "Agente de Retenção B2B: atendimento a clientes empresariais, orientação e negociação orientada por metas, escuta ativa e contorno de objeções (pacote Claro NET)",
                ]},
             ]},
            {"company": "MG VIDROS — AUTO GLASS  |  03/2013 – 09/2015",
             "roles": [
                {"role": "Jovem Aprendiz / Auxiliar de Vendedor Técnico", "bullets": [
                    "Contrato de Jovem Aprendiz com formação profissionalizante de Assistente de Vendas pelo ISBET; atendimento, elaboração de orçamentos e apoio à venda técnica no setor automotivo"
                ]},
             ]},
        ]},
        {"title": "FORMAÇÃO", "bullets": [
            "Administração — Faculdade FATENP — Nova Palhoça (Transferido)",
            "Análise e Desenvolvimento de Sistemas (ADS) — UNISUL — Pedra Branca (cursando, conclusão prevista 2027)",
            "Assistente Administrativo — S.O.S Educação Profissional (2012, 18 meses)",
        ]},
        {"title": "IDIOMAS", "bullets": [
            "Português nativo  |  Inglês avançado",
        ]},
    ],
}


def build_pdf():
    from reportlab.lib.colors import HexColor
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (BaseDocTemplate, Frame, HRFlowable, KeepTogether,
                                    PageTemplate, Paragraph, Spacer)

    ink, muted, accent, rule = (HexColor("#" + c) for c in (INK, MUTED, ACCENT, RULE))
    # Compactação solicitada pelo Rafael: uma página, sem remover fatos.
    s_name = ParagraphStyle("Name", fontName="Helvetica-Bold", fontSize=16, leading=18, textColor=ink, spaceAfter=1)
    s_role = ParagraphStyle("Role", fontName="Helvetica-Bold", fontSize=9, leading=10.5, textColor=ink, spaceAfter=2)
    s_contact = ParagraphStyle("Contact", fontName="Helvetica", fontSize=7.5, leading=9, textColor=muted, spaceAfter=1)
    s_sec = ParagraphStyle("Section", fontName="Helvetica-Bold", fontSize=9.5, leading=10.5, textColor=accent, spaceBefore=5, spaceAfter=1)
    s_co = ParagraphStyle("Company", fontName="Helvetica-Bold", fontSize=8.8, leading=10, textColor=ink, spaceBefore=4, spaceAfter=0.5)
    s_job = ParagraphStyle("Job", fontName="Helvetica-BoldOblique", fontSize=8.8, leading=10, textColor=ink, spaceAfter=0.5)
    s_body = ParagraphStyle("Body", fontName="Helvetica", fontSize=8.5, leading=9.8, textColor=ink, spaceAfter=0.5)
    s_bul = ParagraphStyle("Bullet", parent=s_body, leftIndent=9, firstLineIndent=0, spaceAfter=0,
                           bulletIndent=2)

    _TOTAL = {"n": 0}

    def footer(canv, doc):
        canv.saveState()
        canv.setFont("Helvetica", 6.5)
        canv.setFillColor(muted)
        total = _TOTAL["n"] or doc.page
        canv.drawCentredString(A4[0] / 2, 6 * mm, f"Página {doc.page} de {total}")
        canv.restoreState()

    def _build(out, story):
        doc = BaseDocTemplate(out, pagesize=A4,
                              leftMargin=12 * mm, rightMargin=12 * mm,
                              topMargin=10 * mm, bottomMargin=10 * mm,
                              title="Currículo - Caio Rafael Barros do Nascimento - Analista Administrativo de Vendas",
                              author="Hermes Career Agent")
        frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="main")
        doc.addPageTemplates([PageTemplate(id="p", frames=[frame], onPage=footer)])
        doc.build(story)

    def _make_story():
        st = [Paragraph(CONTENT["name"], s_name),
              Paragraph(CONTENT["role"], s_role),
              HRFlowable(width="100%", thickness=1.2, color=accent, spaceAfter=4),
              Paragraph(CONTENT["contact"], s_contact)]
        for sec in CONTENT["sections"]:
            section_head = [
                Paragraph(sec["title"], s_sec),
                HRFlowable(width="100%", thickness=0.6, color=rule, spaceAfter=4),
            ]
            block = []
            for b in sec.get("body", []):
                block.append(Paragraph(b, s_body))
            for b in sec.get("bullets", []):
                block.append(Paragraph(f"- {b}", s_bul, bulletText=""))
            for job in sec.get("jobs", []):
                jb = [Paragraph(job["company"], s_co)]
                for r in job["roles"]:
                    if r["role"]:
                        jb.append(Paragraph(r["role"], s_job))
                    for b in r["bullets"]:
                        jb.append(Paragraph(f"- {b}", s_bul, bulletText=""))
                block.extend(jb)
            if sec["title"] == "EXPERIÊNCIA PROFISSIONAL":
                # As tarefas podem fluir sem um bloco indivisível, aproveitando uma página.
                st.extend(section_head)
                st.extend(block)
            elif sec["title"] == "FORMAÇÃO":
                st.extend(section_head)
                st.append(KeepTogether(block))
            else:
                # Idiomas pode acompanhar a formação sem ser empurrado para uma página vazia.
                st.extend(section_head)
                st.extend(block)
        return st

    _build(PDF_OUT + ".pass1.pdf", _make_story())
    from pypdf import PdfReader
    _TOTAL["n"] = len(PdfReader(PDF_OUT + ".pass1.pdf").pages)
    _build(PDF_OUT, _make_story())
    try: os.remove(PDF_OUT + ".pass1.pdf")
    except OSError: pass
    return PDF_OUT


def _docx_border_bottom(paragraph, color="C9CFD6", sz="6"):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "1"); bottom.set(qn("w:color"), color)
    pbdr.append(bottom); pPr.append(pbdr)


def _docx_footer_numbers(doc):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    sec = doc.sections[0]
    fp = sec.footer.paragraphs[0]
    fp.alignment = 2  # center=2? 1=center; usar 2=right? -> 1 center
    fp.alignment = 1
    fp.text = ""
    r = fp.add_run(); r.font.size = __import__("docx").shared.Pt(6.5); r.font.color.rgb = __import__("docx").shared.RGBColor(0x5A, 0x62, 0x70)
    def field(run, instr):
        f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
        f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
        run._r.append(f1); run._r.append(it); run._r.append(f2)
    r2 = fp.add_run("Página "); r2.font.size = __import__("docx").shared.Pt(6.5)
    r2.font.color.rgb = __import__("docx").shared.RGBColor(0x5A, 0x62, 0x70)
    field(fp.add_run(), "PAGE")
    r3 = fp.add_run(" de "); r3.font.size = __import__("docx").shared.Pt(6.5)
    r3.font.color.rgb = __import__("docx").shared.RGBColor(0x5A, 0x62, 0x70)
    field(fp.add_run(), "NUMPAGES")


def build_docx():
    from docx import Document
    from docx.shared import Pt, Mm, RGBColor
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    doc = Document()
    for sec in doc.sections:
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin, sec.right_margin = Mm(12), Mm(12)
        sec.top_margin, sec.bottom_margin = Mm(10), Mm(10)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"; style.font.size = Pt(8.5)
    style.font.color.rgb = RGBColor(0x23, 0x27, 0x2F)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.widow_control = True

    def para(text, bold=False, italic=False, size=8.5, color=(0x23, 0x27, 0x2F), before=0, after=0, align=None):
        p = doc.add_paragraph()
        run = p.add_run(text); run.bold = bold; run.italic = italic
        run.font.size = Pt(size); run.font.color.rgb = RGBColor(*color)
        p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
        if align is not None: p.alignment = align
        return p

    para(CONTENT["name"], bold=True, size=16, after=1)
    p = para(CONTENT["role"], bold=True, size=9, after=2)
    _docx_border_bottom(p, color=ACCENT, sz="9")
    para(CONTENT["contact"], size=7.5, color=(0x5A, 0x62, 0x70), after=1)
    for sec in CONTENT["sections"]:
        p = para(sec["title"], bold=True, size=9.5, color=(0x1F, 0x5C, 0x4E), before=4, after=1)
        _docx_border_bottom(p, color=RULE)
        for b in sec.get("body", []):
            para(b)
        for b in sec.get("bullets", []):
            bp = para(f"- {b}")
            bp.paragraph_format.left_indent = Mm(3)
        for job in sec.get("jobs", []):
            para(job["company"], bold=True, size=8.8, before=3, after=0)
            for r in job["roles"]:
                if r["role"]:
                    para(r["role"], bold=True, italic=True, size=8.8, after=0)
                for b in r["bullets"]:
                    bp = para(f"- {b}")
                    bp.paragraph_format.left_indent = Mm(3)
    _docx_footer_numbers(doc)
    doc.save(DOCX_OUT)
    return DOCX_OUT


def write_spec():
    spec = {"title": "Currículo - Caio Rafael Barros do Nascimento - Analista Administrativo de Vendas (V2)",
            "author": "Hermes Career Agent", "version": "V2", "page_size": "A4",
            "design_tokens": {"INK": INK, "MUTED": MUTED, "ACCENT": ACCENT, "RULE": RULE},
            "type_scale": "skill resume-document-design, secao 2 (base V2)",
            "content": deepcopy(CONTENT)}
    json.dump(spec, open(SPEC_OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    return SPEC_OUT


if __name__ == "__main__":
    print("PDF:", build_pdf())
    print("DOCX:", build_docx())
    print("SPEC:", write_spec())
